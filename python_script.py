import sys
import os
import zipfile
import logging
import sqlite3
from datetime import datetime
from dateutil.relativedelta import relativedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

# Import mock tables dictionary from mock_source_tables.py
import mock_source_tables


## 1. Logger Setup
log_dir = os.path.dirname(os.path.abspath(__file__))
log_path = os.path.join(log_dir, "logger_file.log")

logger = logging.getLogger('logger_setup')
logger.setLevel(logging.DEBUG)

formatter = logging.Formatter("%(asctime)s - %(name)s - %(levelname)s - %(message)s")

# File Handler
file_handler = logging.FileHandler(log_path, mode="a", encoding="utf-8")
file_handler.setLevel(logging.DEBUG)
file_handler.setFormatter(formatter)

# Console Handler
console_handler = logging.StreamHandler(sys.stdout)
console_handler.setLevel(logging.DEBUG)
console_handler.setFormatter(formatter)

if not logger.handlers:
    logger.addHandler(file_handler)
    logger.addHandler(console_handler)


## 2. In-Memory SQLite Setup & Data Loading
conn = sqlite3.connect(":memory:")

# Populate mock DataFrames into in-memory SQLite tables
for table_name, df in mock_source_tables.TABLES.items():
    sqlite_table = table_name.replace(".", "__")
    df.to_sql(sqlite_table, conn, index=False, if_exists="replace")


def fetch_local_sqlite(query):
    """Simple wrapper to fetch data from SQLite into pandas DataFrame."""
    return pd.read_sql_query(query, conn)


## 3. Parameter Parsing
param1 = "previous_snapshot_date=2026-07-31"
param2 = "current_snapshot_date=2026-08-31"
param3 = "previous_schema_kb=findata_kb_31_07_2026"
param4 = "current_schema_kb=findata_kb_31_08_2026"
param5 = "previous_schema_dbl=findata_dbl_lap_31_07_2026"
param6 = "current_schema_dbl=findata_dbl_lap_31_08_2026"
param7 = "current_month_start=2026-08-01"
param8 = "current_schema_student=findata_stu_31_08_2026"
param9 = "current_schema_kz=findata_kz_31_08_2026"

previous_snapshot_date = param1.split('=')[1]
current_snapshot_date = param2.split('=')[1]
previous_schema_kb = param3.split('=')[1]
current_schema_kb = param4.split('=')[1]
previous_schema_dbl = param5.split('=')[1]
current_schema_dbl = param6.split('=')[1]
current_month_start = param7.split('=')[1]
current_schema_student = param8.split('=')[1]
current_schema_kz = param9.split('=')[1]

dt = current_snapshot_date[:7]
now = datetime.now()
logger.info("current timing is {}".format(now))
previous_snapshot_start = datetime.strptime(current_month_start, "%Y-%m-%d") - relativedelta(months=1)
previous_snapshot_start = datetime.strftime(previous_snapshot_start, "%Y-%m-%d")

# Output directory and file configurations
output_dir = Path(log_dir) / "Excell_output" / "finance_monthly_segment_analysis"
output_dir.mkdir(parents=True, exist_ok=True)
logger.info("output dir is {}".format(output_dir))

file_path_1 = output_dir / f"KBNBFC({dt}).xlsx"
logger.info("file_path_1 = {}".format(file_path_1))

# Fixed: Create ZIP file directly inside output_dir folder
output_dir_path = output_dir / f"segment_analysis_{dt}.zip"
logger.info("output_dir_path = {}".format(output_dir_path))

file_name = os.path.basename(output_dir_path)
logger.info("File Name is {}".format(file_name))

## 4. SQL Queries

# Query 01: Previous Month POS
previous_month_pos_kb = '''
SELECT
    Platform,
    snapshot_date,
    erp_product_name,
    COALESCE(Platform, '') || COALESCE(erp_product_name, '') AS concat,
    kbnbfc_pos
FROM (
    SELECT
        'KB' AS Platform,
        strftime('%Y-%m', snapshot_date) AS snapshot_date,
        erp_product_name,
        CAST(SUM(kbnbfc_pos) AS REAL) AS kbnbfc_pos
    FROM summary_kb__AUM_loan_snapshot_2_res_3 aum
    JOIN {1}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
    WHERE snapshot_date = '{0}' AND aum.product_name <> 'CARD'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT
        'KB' AS Platform,
        strftime('%Y-%m', a.snapshot_date) AS snapshot_date,
        'CARD' AS erp_product_name,
        CAST(SUM(kbnbfc_pos) AS REAL) AS kbnbfc_pos
    FROM summary_kb__AUM_loan_snapshot_2_res_3 a
    JOIN {1}__yp_loan_dimensional_tbl AS t3 ON (a.purchaseid = t3.purchase_id)
    WHERE a.snapshot_date = '{0}' AND a.product_name = 'CARD'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT
        'KZ' AS Platform,
        strftime('%Y-%m', snapshot_date) AS snapshot_date,
        'KZ' AS erp_product_name,
        CAST(SUM(pos) AS REAL) AS kbnbfc_pos
    FROM summary_kz__AUM_loan_snapshot_2_res_3 aum
    WHERE snapshot_date = '{0}' AND fund_name IN ('KB','KBP')
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'Offline' AS Platform,
        '' AS snapshot_date,
        'Offline' AS erp_product_name,
        CAST(SUM(principal) AS REAL) AS kbnbfc_pos
    FROM {1}__offline_loandata_tbl 
    WHERE paid_date IS NULL
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'DBL_LAP_TWL_POPSA_SCF' AS Platform,
        strftime('%Y-%m', snapshot_date) AS snapshot_date,
        ldt.erp_product_name AS erp_product_name,
        CAST(SUM(kbnbfc_pos) AS REAL) AS kbnbfc_pos
    FROM summary_kb__DBL_LAP_offline_aum_loan_snapshot aum
    LEFT JOIN {2}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
    WHERE snapshot_date = '{0}' AND platform <> 'Offline'
    GROUP BY 1, 2, 3
);'''.format(previous_snapshot_date, previous_schema_kb, previous_schema_dbl)

logger.info('previous_month_pos_kb_query = {}'.format(previous_month_pos_kb))


# Query 02: Current Month POS
current_month_pos_kb = '''
SELECT
    Platform,
    snapshot_date,
    erp_product_name,
    COALESCE(Platform, '') || COALESCE(erp_product_name, '') AS concat,
    kbnbfc_pos,
    NULL AS previous_month_pos,
    NULL AS current_month_disbursement,
    NULL AS current_month_collections,
    NULL AS calculation,
    NULL AS difference
FROM (
    SELECT
        'KB' AS Platform,
        strftime('%Y-%m', snapshot_date) AS snapshot_date,
        erp_product_name,
        CAST(SUM(kbnbfc_pos) AS REAL) AS kbnbfc_pos
    FROM summary_kb__AUM_loan_snapshot_2_res_3 aum
    JOIN {1}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
    WHERE snapshot_date = '{0}' AND aum.product_name <> 'CARD'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT
        'KB' AS Platform,
        strftime('%Y-%m', a.snapshot_date) AS snapshot_date,
        'CARD' AS erp_product_name,
        CAST(SUM(kbnbfc_pos) AS REAL) AS kbnbfc_pos
    FROM summary_kb__AUM_loan_snapshot_2_res_3 a
    JOIN {1}__yp_loan_dimensional_tbl AS t3 ON a.purchaseid = t3.purchase_id
    WHERE a.snapshot_date = '{0}' AND a.product_name = 'CARD'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT
        'KZ' AS Platform,
        strftime('%Y-%m', snapshot_date) AS snapshot_date,
        'KZ' AS erp_product_name,
        CAST(SUM(pos) AS REAL) AS kbnbfc_pos
    FROM summary_kz__AUM_loan_snapshot_2_res_3 aum
    WHERE snapshot_date = '{0}' AND fund_name IN ('KB','KBP')
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT
        'Offline' AS Platform,
        '' AS snapshot_date,
        'Offline' AS erp_product_name,
        CAST(SUM(principal) AS REAL) AS kbnbfc_pos
    FROM {1}__offline_loandata_tbl
    WHERE paid_date IS NULL
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT
        'DBL_LAP_TWL_POPSA_SCF' AS Platform,
        strftime('%Y-%m', snapshot_date) AS snapshot_date,
        ldt.erp_product_name AS erp_product_name,
        CAST(SUM(kbnbfc_pos) AS REAL) AS kbnbfc_pos
    FROM summary_kb__DBL_LAP_offline_aum_loan_snapshot aum
    LEFT JOIN {2}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
    WHERE snapshot_date = '{0}' AND platform <> 'Offline'
    GROUP BY 1, 2, 3
);'''.format(current_snapshot_date, current_schema_kb, current_schema_dbl)

logger.info('current_month_pos_kb_query = {}'.format(current_month_pos_kb))


# Query 03: Current Month Disbursement
current_month_disbursement_micro_lap = '''
SELECT
    Platform,
    disb_mon,
    Product_name,
    COALESCE(Platform, '') || COALESCE(product_name, '') AS concat,
    loan_count,
    amount_disb
FROM (
    SELECT 
        'KB' AS Platform,
        DATE(disbursed_date) AS disb_mon,
        erp_product_name AS product_name,
        COUNT(DISTINCT emi.loan_id) AS loan_count,
        CAST(SUM(CASE WHEN installment_number = 1 THEN kbnbfc_principaldue END) AS REAL) AS amount_disb
    FROM {2}__yp_emi_data_tbl emi
    LEFT JOIN {2}__yp_loan_dimensional_tbl ldt ON emi.loan_id = ldt.loan_id
    WHERE DATE(disbursed_date) BETWEEN '{0}' AND '{1}'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'Offline' AS Platform,
        DATE(disbursed_date) AS disb_mon,
        'Offline' AS product_name,
        COUNT(DISTINCT Curr_New_Account_No) AS loan_count,
        CAST(SUM(principal) AS REAL) AS amount_disb
    FROM {2}__offline_loandata_tbl
    WHERE DATE(disbursed_date) BETWEEN '{0}' AND '{1}'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'DBL_LAP_TWL_POPSA_SCF' AS Platform,
        DATE(disbursed_date) AS disb_mon,
        ldt.erp_product_name AS product_name,
        COUNT(DISTINCT emi.loan_id) AS loan_count,
        CAST(SUM(CASE WHEN installment_number = 1 THEN principaldue END) AS REAL) AS amount_disb
    FROM {3}__yp_emi_data_tbl emi
    LEFT JOIN {3}__yp_loan_dimensional_tbl ldt ON emi.loan_id = ldt.loan_id
    WHERE DATE(disbursed_date) BETWEEN '{0}' AND '{1}'
      AND ldt.erp_product_name NOT IN ('LAP','MICRO_LAP')
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'DBL_LAP_TWL_POPSA_SCF' AS Platform,
        DATE(datetime(b.disbursedon, '+5 hours', '+30 minutes')) AS disb_mon,
        c.erp_product_name AS product_name,
        COUNT(a.kbloanid) AS loan_count,
        CAST(SUM(CASE
            WHEN trancheNo = 1 THEN (CAST(amount AS REAL) + CAST(agreementFees AS REAL) +
                                     CAST(processingFees AS REAL) + CAST(insuranceFee AS REAL) +
                                     CAST(gstFees AS REAL))
            ELSE CAST(amount AS REAL)
        END) AS REAL) AS amount_disb
    FROM {3}__yp_loan a
    JOIN {3}__yp_tranche_disbursals b ON a.id = b.loanid
    LEFT JOIN {3}__yp_loan_dimensional_tbl c ON a.id = c.loan_id
    WHERE datetime(b.disbursedon, '+5 hours', '+30 minutes') >= '{0} 00:00:00'
      AND datetime(b.disbursedon, '+5 hours', '+30 minutes') <= '{1} 23:59:59'
      AND b.state = 87 AND c.erp_product_name IN ('LAP','MICRO_LAP')
    GROUP BY 1, 2, 3
    ORDER BY 1, 3, 2 DESC
);'''.format(current_month_start, current_snapshot_date, current_schema_kb, current_schema_dbl)

logger.info('current_month_disbursement_micro_lap = {}'.format(current_month_disbursement_micro_lap))


# Query 04: Current Month Collections
current_month_collections_kb = '''
SELECT
    Platform,
    collected_month,
    erp_product_name,
    COALESCE(Platform,'') || COALESCE(erp_product_name,'') AS concat,
    principal_collection_kbnbfc
FROM (
    SELECT 
        'KB' AS Platform,
        strftime('%Y-%m-%d', txndate) AS collected_month,
        erp_product_name,
        CAST(SUM(COALESCE(kbnbfc_principal_paid, 0)) AS REAL) AS principal_collection_kbnbfc
    FROM {2}__yp_collections_data_tbl col
    LEFT JOIN {2}__yp_loan_dimensional_tbl ldt ON col.loanid = ldt.loan_id
    WHERE DATE(txndate) BETWEEN '{0}' AND '{1}'
      AND ldt.product_name != 'CARD'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'KB' AS Platform,
        strftime('%Y-%m-%d', principal_paid_date) AS collected_month,
        'CARD' AS erp_product_name,
        CAST(SUM(principal_paid) AS REAL) AS principal_collection_kbnbfc
    FROM {2}__card_new_recon
    WHERE DATE(principal_paid_date) BETWEEN '{0}' AND '{1}'
      AND purchase_type IN ('purchase', 'bankWithdrawal')
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'KZ' AS Platform,
        strftime('%Y-%m-%d', emi_partital_paid_date) AS collected_month,
        'KZ' AS erp_product_name,
        CAST(SUM(principal_paid) AS REAL) AS principal_collection_kbnbfc
    FROM {4}__yp_emi_data_tbl a
    WHERE emi_partital_paid_date BETWEEN '{0} 00:00:00' AND '{1} 23:59:59'
      AND fund_name IN ('KB', 'KBP')
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'offline' AS Platform,
        strftime('%Y-%m-%d', paid_date) AS collected_month,
        'offline' AS erp_product_name,
        CAST(SUM(principalpaid) AS REAL) AS principal_collection_kbnbfc
    FROM {2}__offline_loandata_tbl
    WHERE DATE(paid_date) BETWEEN '{0}' AND '{1}'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'DBL_LAP_TWL_POPSA_SCF' AS Platform,
        strftime('%Y-%m-%d', txndate) AS collected_month,
        ldt.erp_product_name,
        CAST(SUM(COALESCE(principal_paid, 0)) AS REAL) AS principal_collection_kbnbfc
    FROM {3}__yp_collections_data_tbl col
    LEFT JOIN {3}__yp_loan_dimensional_tbl ldt ON col.loanid = ldt.loan_id
    WHERE DATE(txndate) BETWEEN '{0}' AND '{1}'
    GROUP BY 1, 2, 3

    UNION ALL

    SELECT 
        'Krazybee_Student' AS Platform,
        strftime('%Y-%m-%d', paiddate) AS collected_month,
        'Krazybee_Student' AS erp_product_name,
        CAST(SUM(PrincipalPaid) AS REAL) AS principal_collection_kbnbfc
    FROM {5}__temp_all_emi_all_vw
    WHERE paiddate BETWEEN '{0} 00:00:00' AND '{1} 23:59:59'
      AND fund_name = 'KBNBFC'
    GROUP BY 1, 2

    ORDER BY 1, 3, 2 DESC
);'''.format(current_month_start, current_snapshot_date, current_schema_kb, current_schema_dbl, current_schema_kz, current_schema_student)

logger.info('current_month_collections_kb = {}'.format(current_month_collections_kb))


# Query 05: Partner Previous Month POS
partner_previous_month_pos_kb = '''
SELECT 
    strftime('%Y-%m', snapshot_date) AS snapshot_date,
    erp_product_name,
    CAST(SUM(partner_pos) AS REAL) AS partner_pos
FROM summary_kb__AUM_loan_snapshot_2_res_3 aum
JOIN {1}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
WHERE snapshot_date = '{0}' AND aum.partner_pos > 0
GROUP BY 1, 2
ORDER BY 1, 2;'''.format(previous_snapshot_date, previous_schema_kb)

logger.info('partner_previous_month_pos_kb = {}'.format(partner_previous_month_pos_kb))


# Query 06: Partner Current Month POS
partner_pos_current_month_kb = '''
SELECT 
    strftime('%Y-%m', snapshot_date) AS snapshot_date,
    erp_product_name,
    CAST(SUM(partner_pos) AS REAL) AS partner_pos,
    NULL AS previous_pos,
    NULL AS current_disbursement,
    NULL AS current_collections,
    NULL AS calculations,
    NULL AS difference
FROM summary_kb__AUM_loan_snapshot_2_res_3 aum
JOIN {1}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
WHERE snapshot_date = '{0}' AND aum.partner_pos > 0
GROUP BY 1, 2
ORDER BY 1, 2;'''.format(current_snapshot_date, current_schema_kb)

logger.info('partner_pos_current_month_kb = {}'.format(partner_pos_current_month_kb))


# Query 07: Partner Current Disbursement
partner_current_disbursement_kb = '''
SELECT
    strftime('%Y-%m-%d', disbursed_date) AS disb_mon,
    erp_product_name,
    COUNT(DISTINCT emi.loan_id) AS loan_count,
    CAST(SUM(CASE WHEN installment_number = 1 THEN partner_principaldue END) AS REAL) AS amount_disb
FROM {2}__yp_emi_data_tbl emi
JOIN {2}__yp_loan_dimensional_tbl ldt ON emi.loan_id = ldt.loan_id
WHERE DATE(disbursed_date) BETWEEN '{0}' AND '{1}'
  AND emi.state IN ('disbursed', 'closed')
  AND emi.kbnbfc_fund_allocation_percent BETWEEN 1 AND 99
GROUP BY 1, 2
ORDER BY 1 DESC, 2;'''.format(current_month_start, current_snapshot_date, current_schema_kb)

logger.info('partner_current_disbursement_kb = {}'.format(partner_current_disbursement_kb))


# Query 08: Partner Collections
partner_current_month_collections_kb = '''
SELECT
    strftime('%Y-%m-%d', txndate) AS collected_month,
    erp_product_name,
    CAST(SUM(COALESCE(partner_principal_paid, 0)) AS REAL) AS principal_collection_Partner
FROM {2}__yp_collections_data_tbl col
LEFT JOIN {2}__yp_loan_dimensional_tbl ldt ON col.loanid = ldt.loan_id
WHERE DATE(txndate) BETWEEN '{0}' AND '{1}'
  AND partner_principal_paid > 0
GROUP BY 1, 2
ORDER BY 1 DESC;'''.format(current_month_start, current_snapshot_date, current_schema_kb)

logger.info('partner_current_month_collections_kb = {}'.format(partner_current_month_collections_kb))


# Query 09: Installment Count
Installment_count_kb = '''
SELECT 
    'KB' AS Platform,
    ldt.erp_product_name,
    COUNT(DISTINCT ldt.loan_id) AS Loan_count,
    COUNT(emi.id) AS Installment_count
FROM {0}__yp_loan_dimensional_tbl ldt
JOIN {0}__yp_loan_installments emi ON emi.loanid = ldt.loan_id
GROUP BY 1, 2

UNION ALL

SELECT 
    'KZ' AS Platform,
    l.productname,
    COUNT(DISTINCT l.id) AS Loan_Count,
    COUNT(li.id) AS Installment_count
FROM {1}__yp_loan l
JOIN {1}__yp_loan_installments li ON l.id = li.loanid
GROUP BY 1, 2

UNION ALL

SELECT 
    'DBL_LAP_TWL_POPSA_SCF' AS Platform,
    ldt.product_name,
    COUNT(DISTINCT emi.loanid) AS loan_count,
    COUNT(emi.id) AS Installment_count
FROM {2}__yp_loan_installments emi
LEFT JOIN {2}__yp_loan_dimensional_tbl ldt ON emi.loanid = ldt.loan_id
GROUP BY 1, 2

UNION ALL

SELECT 
    'Offline' AS Platform,
    '' AS product_name,
    COUNT(DISTINCT Curr_New_Account_No) AS loan_count,
    COUNT(*) AS installment_count
FROM {0}__offline_loandata_tbl
GROUP BY 1, 2

UNION ALL

SELECT 
    'Krazybee-Stu' AS Platform,
    '' AS product_name,
    COUNT(DISTINCT OrderID) AS loan_count,
    COUNT(*) AS installment_count
FROM {3}__temp_all_emi_all_vw
GROUP BY 1, 2

ORDER BY 1, 2;'''.format(current_schema_kb, current_schema_kz, current_schema_dbl, current_schema_student)

logger.info('Installment_count_kb = {}'.format(Installment_count_kb))


# Query 10: Daily AUM KB
daily_aum_kb = '''
SELECT
    snapshot_date,
    ldt.erp_product_name AS product,
    SUM(pos) AS pos,
    SUM(pan_level_aum_kbnbfc) AS total_aum,
    SUM(da_assignee_pan_level_aum) AS da_aum,
    SUM(pan_level_aum_kbnbfc) - SUM(da_assignee_pan_level_aum) AS net_aum
FROM kreditbee_bi_dw_athena__kb_aum_loan_snapshot_archived aum
LEFT JOIN {0}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
WHERE snapshot_date BETWEEN '{1}' AND '{2}'
GROUP BY 1, 2
ORDER BY 1;'''.format(current_schema_kb, current_month_start, current_snapshot_date)

logger.info('daily_aum_kb = {}'.format(daily_aum_kb))


# Query 11: Daily AUM Offline
daily_aum_offline = '''
SELECT
    snapshot_date,
    ldt.erp_product_name AS product,
    SUM(pos) AS pos,
    SUM(pan_level_aum_kbnbfc) AS total_aum
FROM kreditbee_bi_dw_athena__offline_aum_loan_snapshot_archived aum
LEFT JOIN {0}__yp_loan_dimensional_tbl ldt ON aum.loan_id = ldt.loan_id
WHERE snapshot_date BETWEEN '{1}' AND '{2}'
GROUP BY 1, 2
ORDER BY 1;'''.format(current_schema_dbl, current_month_start, current_snapshot_date)

logger.info('daily_aum_offline = {}'.format(daily_aum_offline))


# Query 12: Krazybee Student
krazybee_student = '''
SELECT
    strftime('%Y-%m-%d', paiddate) AS paid_mon,
    'Krazybee' AS Product,
    CAST(SUM(PrincipalPaid) AS REAL) AS principal_collection_kbnbfc
FROM {0}__temp_all_emi_all_vw
WHERE paiddate BETWEEN '{1} 00:00:00' AND '{2} 23:59:59' 
  AND fund_name = 'KBNBFC'
GROUP BY 1, 2
ORDER BY 1 DESC;'''.format(current_schema_student, previous_snapshot_start, current_snapshot_date)

logger.info('krazybee_student = {}'.format(krazybee_student))


## 5. Pipeline Execution
def queries_for_kb():
    try:
        logger.info("Executing local queries using in-memory SQLite...")

        df1 = fetch_local_sqlite(previous_month_pos_kb)
        logger.info("previous_month_pos_kb data fetched successfully moving towards next...")
        
        df2 = fetch_local_sqlite(current_month_pos_kb)
        logger.info("current_month_pos_kb data fetched successfully moving towards next...")

        df3 = fetch_local_sqlite(current_month_disbursement_micro_lap)
        logger.info("current_month_disbursement_micro_lap data fetched successfully moving towards next...")

        df4 = fetch_local_sqlite(current_month_collections_kb)
        logger.info("current_month_collections_kb data fetched successfully moving towards next...")
       
        df5 = fetch_local_sqlite(partner_previous_month_pos_kb)
        logger.info("partner_previous_month_pos_kb data fetched successfully moving towards next...")

        df6 = fetch_local_sqlite(partner_pos_current_month_kb)
        logger.info("partner_pos_current_month_kb data fetched successfully moving towards next...")

        df7 = fetch_local_sqlite(partner_current_disbursement_kb)
        logger.info("partner_current_disbursement_kb data fetched successfully moving towards next...")

        df8 = fetch_local_sqlite(partner_current_month_collections_kb)
        logger.info("partner_current_month_collections_kb data fetched successfully moving towards next...")

        df9 = fetch_local_sqlite(Installment_count_kb)
        logger.info("Installment_count_kb data fetched successfully moving towards next...")

        df10 = fetch_local_sqlite(daily_aum_kb)
        logger.info("daily_aum_kb data fetched successfully moving towards next...")

        df11 = fetch_local_sqlite(daily_aum_offline)
        logger.info("daily_aum_offline data fetched successfully moving towards next...")

        df12 = fetch_local_sqlite(krazybee_student)
        logger.info("krazybee_student data fetched successfully moving towards next...")

        logger.info("All DataFrames fetched successfully from SQLite. Writing to Excel...")

        with pd.ExcelWriter(file_path_1, engine='openpyxl') as writer:
            df1.to_excel(writer, sheet_name='KBNBFC POS', startcol=0, index=False)
            df2.to_excel(writer, sheet_name='KBNBFC POS', startcol=8, index=False)

            df3.to_excel(writer, sheet_name='KBNBFC Disb coll', startcol=0, index=False)
            df4.to_excel(writer, sheet_name='KBNBFC Disb coll', startcol=9, index=False)

            df5.to_excel(writer, sheet_name='partner disb coll', startcol=0, index=False)
            df6.to_excel(writer, sheet_name='partner disb coll', startcol=6, index=False)
            df7.to_excel(writer, sheet_name='partner disb coll', startcol=15, index=False)
            df8.to_excel(writer, sheet_name='partner disb coll', startcol=21, index=False)

            df9.to_excel(writer, sheet_name='Installment count', startcol=0, index=False)
            df10.to_excel(writer, sheet_name='daily aum kb', startcol=0, index=False)
            df11.to_excel(writer, sheet_name='daily aum offline', startcol=0, index=False)
            df12.to_excel(writer, sheet_name='Krazybee Student', startcol=0, index=False)

            # Styling Excel Output
            header_font = Font(name='Calibri', size=11, bold=True, color='002060')
            header_fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
            thin_border = Border(
                left=Side(style='thin', color='A6A6A6'),
                right=Side(style='thin', color='A6A6A6'),
                top=Side(style='thin', color='A6A6A6'),
                bottom=Side(style='thin', color='A6A6A6')
            )
            data_border = Border(
                left=Side(style='thin', color='D3D3D3'),
                right=Side(style='thin', color='D3D3D3'),
                top=Side(style='thin', color='D3D3D3'),
                bottom=Side(style='thin', color='D3D3D3')
            )

            for sheet_name in writer.sheets:
                ws = writer.sheets[sheet_name]
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    is_header_row = (row[0].row == 1)
                    for cell in row:
                        if is_header_row:
                            if cell.value is not None:
                                cell.font = header_font
                                cell.fill = header_fill
                                cell.alignment = Alignment(horizontal='center', vertical='center')
                            cell.border = thin_border
                        else:
                            cell.border = data_border
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = col[0].column_letter
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        logger.info(f"Segment Analysis Excel generated successfully: {file_path_1}")
        return [str(file_path_1)]

    except Exception as e:
        logger.error(f"Exception occurred in queries_for_kb: {e}", exc_info=True)
        return []


def copy_and_zip_file(file_paths):
    try:
        with zipfile.ZipFile(output_dir_path, 'w') as zipf:
            for file in file_paths:
                if os.path.exists(file):
                    zipf.write(file, os.path.basename(file))
                    logger.info(f"Added to ZIP: {file}")
                else:
                    logger.warning(f"File not found: {file}")
        
        logger.info(f"Created ZIP file: {output_dir_path}")

        # Remove the standalone original Excel file(s) outside the zip
        for file in file_paths:
            if os.path.exists(file):
                os.remove(file)
                logger.info(f"Cleaned up standalone file: {file}")

        print("SUCCESS")
        print(f"Zip created locally at: {output_dir_path}")

    except Exception as e:
        logger.error(f"Exception occurred in copy_and_zip_file: {e}")
        
if __name__ == "__main__":
    generated_files = queries_for_kb()
    if generated_files:
        copy_and_zip_file(generated_files)